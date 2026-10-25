"""One strategy answer to a .docx: runs scripts/dryrun_strategy.py with a flag set (read-only, in its own
process, the live .env plus --set) and writes the full answer and its cited sources -- the articles and the
rulings the app shows under the answer (document, section, text) -- to a Word file.

    cd /opt/chatbot && venv/bin/python scripts/answer_docx.py --text "<the message>" \\
        --set FACT_SIGNALS_ENABLED=true --out /tmp/risposta.docx
    ... --file "tests/po_cases/Leonardo Agricoltore.docx" --request "<the message>"   (a document and a message)
    ... --case h3                                                                     (an eval case)
    ... --from-json saved.json                                                        (a dryrun --save, no run)

--blind leaves the flag set out of the file (scripts/ab_answers.py); --label is the title ("Risposta 1").
"""

import argparse
import json
import os
import re
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
_INLINE_BOLD = re.compile(r"\*\*(.+?)\*\*")


def run_dryrun(args) -> dict:
    """dryrun_strategy's saved record (status, answer, citations) for this source and flag set."""
    with tempfile.TemporaryDirectory() as tmp:
        saved = os.path.join(tmp, "run.json")
        cmd = [sys.executable, os.path.join(HERE, "dryrun_strategy.py"), "--save", saved, "--env", args.env,
               "--lang", args.lang]
        if args.text:
            cmd += ["--text", args.text]
        elif args.case:
            cmd += ["--case", args.case]
        else:
            cmd += ["--file", args.file] + (["--request", args.request] if args.request else [])
        for item in args.set:
            cmd += ["--set", item]
        done = subprocess.run(cmd, cwd=os.getcwd(), capture_output=True, text=True)
        if not os.path.isfile(saved):
            sys.exit(f"dryrun_strategy wrote nothing (exit {done.returncode}):\n{done.stdout[-2000:]}\n{done.stderr[-2000:]}")
        with open(saved, encoding="utf-8") as fh:
            return json.load(fh)


def _runs(paragraph, text: str) -> None:
    """text into paragraph, **bold** spans in bold."""
    pos = 0
    for m in _INLINE_BOLD.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        paragraph.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_answer(doc, answer: str) -> None:
    """The answer's markdown as Word: "**1. Titolo**" and "#..." lines as headings, "- " lines as bullets."""
    for raw in (answer or "").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        heading = re.fullmatch(r"\s*\*\*(\d+\.\s+.+?)\*\*\s*", line)
        hashes = re.match(r"\s*(#{1,6})\s+(.*)", line)
        if heading:
            doc.add_heading(heading.group(1), level=2)
        elif hashes:
            doc.add_heading(hashes.group(2).strip("* "), level=min(4, len(hashes.group(1)) + 1))
        elif re.match(r"\s*[-•*]\s+", line):
            _runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-•*]\s+", "", line))
        else:
            _runs(doc.add_paragraph(), line.strip())


def _section_label(section: dict, ruling: bool) -> str:
    name, title = str(section.get("name") or ""), str(section.get("title") or "").strip()
    if ruling:
        return title or name
    number = name.split(".")[0] if re.match(r"\d", name) else name
    return f"Art. {number}" + (f" — {title}" if title else "")


def add_sources(doc, citations: list, chars: int = 3000) -> None:
    """The cited sources as the app lists them under the answer: norms, then rulings; each document with
    its sections (label and text)."""
    rulings = [c for c in citations or [] if (c.get("document_type") or "") == "interpretation"]
    norms = [c for c in citations or [] if c not in rulings]
    doc.add_heading("Fonti citate", level=1)
    for title, group, ruling in (("Norme", norms, False), ("Giurisprudenza", rulings, True)):
        if not group:
            continue
        doc.add_heading(title, level=2)
        for c in group:
            p = doc.add_paragraph()
            p.add_run(str(c.get("document_name") or c.get("document_id") or "")).bold = True
            for section in c.get("sections") or []:
                label = _section_label(section, ruling)
                if label:
                    doc.add_paragraph().add_run(label).italic = True
                text = " ".join(str(section.get("plain_text") or "").split())
                if text:
                    doc.add_paragraph(text[:chars] + ("…" if len(text) > chars else ""))


def build(record: dict, out: str, label: str = "Risposta", question: str = "", settings=None) -> str:
    """The .docx of one dryrun record: title, the question, the answer, the sources; the flag set only when
    settings is given (never with --blind)."""
    from docx import Document
    doc = Document()
    doc.core_properties.author = ""
    doc.core_properties.title = label
    doc.add_heading(label, level=0)
    if question:
        doc.add_heading("Domanda", level=1)
        for part in question.split("\n"):
            if part.strip():
                doc.add_paragraph(part.strip())
    doc.add_heading("Risposta", level=1)
    if record.get("status") != "ok":
        doc.add_paragraph(f"(nessuna risposta: {record.get('status')})")
    add_answer(doc, record.get("answer") or record.get("draft") or "")
    add_sources(doc, record.get("citations") or [])
    if settings is not None:
        doc.add_heading("Configurazione", level=1)
        doc.add_paragraph(", ".join(settings) or "il .env, senza modifiche")
    doc.save(out)
    return out


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--text", help="a message with no document")
    source.add_argument("--file", help="a document, read as an upload is (with --request)")
    source.add_argument("--case", help="an eval case id")
    source.add_argument("--from-json", help="a dryrun_strategy --save file: no run")
    parser.add_argument("--request", default="", help="the message sent with --file")
    parser.add_argument("--set", action="append", default=[], metavar="NAME=VALUE")
    parser.add_argument("--out", required=True, help="the .docx to write")
    parser.add_argument("--label", default="Risposta")
    parser.add_argument("--blind", action="store_true", help="no flag set in the file")
    parser.add_argument("--save-json", help="also keep dryrun's record here")
    parser.add_argument("--env", default=os.path.join(os.getcwd(), ".env"))
    parser.add_argument("--lang", default="it")
    args = parser.parse_args()
    if args.from_json:
        with open(args.from_json, encoding="utf-8") as fh:
            record = json.load(fh)
    else:
        record = run_dryrun(args)
    if args.save_json:
        with open(args.save_json, "w", encoding="utf-8") as fh:
            json.dump(record, fh, ensure_ascii=False, indent=1)
    question = args.text or record.get("request") or args.request or ""
    build(record, args.out, args.label, question, None if args.blind else args.set)
    print(f"{args.out}: {record.get('status')}, {len(record.get('answer') or '')} characters, "
          f"{len(record.get('citations') or [])} sources")


if __name__ == "__main__":
    main()
