import json
import re
import shutil
from pathlib import Path
from docx import Document
from openai import OpenAI
import os

CATALOG_PATH = Path('/opt/chatbot/data/system_templates/catalog_enriched.json')
TEMPLATES_DIR = Path('/opt/chatbot/data/system_templates')
NEW_DIR = Path('/opt/chatbot/data/new_templates')
LOG_PATH = Path('/tmp/add_templates.log')

# Load env
from dotenv import load_dotenv
load_dotenv('/opt/chatbot/.env')

client = OpenAI(
    base_url=os.environ['LLM_BASE_URL'],
    api_key=os.environ.get('LLM_API_KEY', 'none'),
)
MODEL = os.environ.get('LLM_MODEL', 'qwen2.5-32b-awq')

def log(msg):
    print(msg)
    with open(LOG_PATH, 'a') as f:
        f.write(msg + '\n')

def extract_text_from_docx(path):
    doc = Document(path)
    return '\n'.join(p.text.strip() for p in doc.paragraphs if p.text.strip())

def extract_sections_from_docx(path):
    doc = Document(path)
    sections = []
    current_heading = None
    current_content = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        is_heading = (
            para.style.name.startswith('Heading') or
            (len(text) < 80 and text.isupper()) or
            (para.runs and all(r.bold for r in para.runs if r.text.strip()))
        )
        if is_heading:
            if current_heading and current_content:
                sections.append({'heading': current_heading, 'content': '\n'.join(current_content).strip()})
            current_heading = text
            current_content = []
        else:
            current_content.append(text)
    if current_heading and current_content:
        sections.append({'heading': current_heading, 'content': '\n'.join(current_content).strip()})
    if not sections:
        all_text = '\n'.join(p.text.strip() for p in doc.paragraphs if p.text.strip())
        sections = [{'heading': 'Contenuto', 'content': all_text}]
    return sections

def slugify(text):
    text = text.lower()
    text = re.sub(r'[àáâã]', 'a', text)
    text = re.sub(r'[èéêë]', 'e', text)
    text = re.sub(r'[ìíîï]', 'i', text)
    text = re.sub(r'[òóôõ]', 'o', text)
    text = re.sub(r'[ùúûü]', 'u', text)
    text = re.sub(r'[^a-z0-9\s-]', '', text)
    text = re.sub(r'[\s]+', '-', text.strip())
    return text

def codice_prefix(codice):
    mapping = {
        'Codice di procedura civile': 'civile',
        'Codice di procedura penale': 'penale',
        'Codice del processo amministrativo': 'amministrativo',
        'Contratti e Atti Stragiudiziali': 'contratti',
        'Diritto del Lavoro': 'lavoro',
        'Arbitrato e Procedure Alternative': 'arbitrato',
        'Diritto Concorsuale': 'concorsuale',
        'Diritto Amministrativo Sanzionatorio': 'sanzionatorio',
    }
    return mapping.get(codice, 'altro')

def get_metadata_from_llm(title, text_preview):
    prompt = f"""Sei un esperto di diritto italiano. Analizza questo documento legale e restituisci SOLO un oggetto JSON valido senza markdown.

Titolo: {title}
Contenuto (prime 1500 caratteri): {text_preview[:1500]}

Restituisci JSON con questi campi:
- "codice": uno tra ["Codice di procedura civile", "Codice di procedura penale", "Codice del processo amministrativo", "Contratti e Atti Stragiudiziali", "Diritto del Lavoro", "Arbitrato e Procedure Alternative", "Diritto Concorsuale", "Diritto Amministrativo Sanzionatorio"]
- "categorie": lista con una categoria descrittiva (es. ["Atti introduttivi", "Misure cautelari", "Atti difensivi"])
- "sottocategorie": lista con una sottocategoria specifica
- "tipo_atto": nome preciso dell'atto (es. "Ricorso per Separazione Giudiziale")
- "label": versione breve del tipo_atto
- "description": una frase che descrive lo scopo dell'atto
- "fields": lista di 3-8 nomi di campo snake_case che l'utente deve compilare (es. ["nome_cliente", "tribunale", "data_udienza"])"""

    response = client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": prompt}],
        max_tokens=600,
        temperature=0.1,
    )
    raw = response.choices[0].message.content.strip()
    raw = re.sub(r'```(?:json)?\s*', '', raw).strip().rstrip('`').strip()
    return json.loads(raw)

def clean_title(filename):
    name = Path(filename).stem
    name = re.sub(r'^\d+\s*[-–]\s*', '', name)
    return name.strip()

# Files that had no match in previous script
unmatched = [
    '001 - Atto di compromesso arbitrale.docx',
    '002 - Clausola per arbitrato irrituale.docx',
    '003 - Clausola per arbitrato rituale.docx',
    '004 - Istanza di esecutorieta del lodo arbitrale.docx',
    '005 - Nomina Arbitro e Domanda di Arbitrato.docx',
    '006 - Ricorso per Decreto Ingiuntivo - Crediti di Lavoro.docx',
    '007 - Decreto Ingiuntivo per Crediti di Lavoro.docx',
    '008 - Diffida per il pagamento del TFR e competenze di fine rapporto.docx',
    '009 - Diffida alla Pubblica Amministrazione.docx',
    '010 - Dimissioni per Giusta Causa.docx',
    '011 - Lettera di Dimissioni Volontarie.docx',
    '012 - Impugnazione Stragiudiziale del Licenziamento.docx',
    '013 - Istanza di Accesso agli Atti - Rapporto di Lavoro.docx',
    '014 - Memoria difensiva avverso ricorso al giudice del lavoro.docx',
    '015 - Opposizione a Licenziamento Disciplinare.docx',
    '016 - Revoca delle Dimissioni.docx',
    '017 - Richiesta Pagamento Ferie Non Godute.docx',
    '018 - Ricorso al Giudice del Lavoro.docx',
    '019 - Ricorso al Giudice del Lavoro - Rito Fornero.docx',
    '020 - Ricorso al TAR - Pubblico Impiego.docx',
    '021 - Risoluzione Consensuale del Rapporto di Lavoro.docx',
    '022 - Tentativo Obbligatorio di Conciliazione.docx',
    '023 - Accordo transattivo (transazione) in materia di lavoro.docx',
    '024 - Verbale di conciliazione in materia di lavoro.docx',
    '025 - Verbale di conciliazione monocratica (Ispettorato Territoriale del Lavoro).docx',
    '032 - Atto di transazione.docx',
    '033 - Contratto locazione commerciale.docx',
    '034 - Contratto di locazione turistica.docx',
    '035 - Facsimile contratto locazione ad uso abitativo.docx',
    '040 - Ricorso in Appello contro Sentenza di Divorzio.docx',
    '042 - Decreto ingiuntivo provvisoriamente esecutivo.docx',
    '044 - Ricorso per Manutenzione nel Possesso.docx',
    '045 - Ricorso per Reintegra nel Possesso.docx',
    '046 - Ricorso per accertamento tecnico preventivo (art. 696 c.p.c.).docx',
    '047 - Ricorso ex art. 696-bis c.p.c.docx',
    '049 - Ricorso per scioglimento del matrimonio cessazione degli effetti civili.docx',
    '051 - Ricorso per Separazione Giudiziale.docx',
    '053 - Atto di Citazione per Usucapione.docx',
    '054 - Atto di citazione per risarcimento danni da sinistro stradale (RCA).docx',
    '057 - Atto di Citazione per Azione Revocatoria Ordinaria.docx',
    '059 - Citazione per Divisione Ereditaria.docx',
    '060 - Atto di Citazione per Revocatoria Fondo Patrimoniale.docx',
    '061 - Citazione Testimoni Processo Civile.docx',
    '062 - Comparsa di Risposta con Chiamata in Causa del Terzo.docx',
    '064 - Formula di conformita su sentenza.docx',
    '065 - Ricorso Procedimento Sommario di Cognizione.docx',
    '066 - Atto di Citazione per Impugnazione Delibera Condominiale.docx',
    '067 - Istanza di conciliazione in sede non contenziosa (art. 322 c.p.c.).docx',
    '068 - Istanza di Correzione Errore Materiale.docx',
    '069 - Istanza di ammissione al patrocinio a spese dello Stato.docx',
    '071 - Preventivo scritto obbligatorio – Fac-simile per prestazioni legali.docx',
    '073 - Procura generale per Trust.docx',
    '074 - Relata di Notifica a mezzo PEC.docx',
    '075 - Istanza di Revoca o Modifica Ordinanza.docx',
    '076 - Ricorso per la nomina dell\'Amministratore di Sostegno.docx',
    '077 - Ricorso Opposizione avverso ingiunzione fiscale (art. 2 R.D. 639 1910).docx',
    '078 - Ricorso Avverso Multe e Violazioni Codice della Strada.docx',
    '079 - Ricorso GDP Specializzato Autovelox Tutor.docx',
    '080 - Ricorso al Giudice di Pace contro Multa.docx',
    '081 - Ricorso al Prefetto contro Verbale di Contravvenzione.docx',
    '083 - Rinuncia agli Atti del Giudizio.docx',
    '086 - Atto di precetto su sentenza civile.docx',
    '093 - Atto di precetto per rilascio di immobile.docx',
    '094 - Atto di precetto su assegno o cambiale.docx',
    '095 - Precetto su sentenza penale di condanna.docx',
    '096 - Verbale di pignoramento mobiliare presso il debitore.docx',
    '097 - Domanda di ammissione al passivo restituzione – rivendicazione.docx',
    '098 - Domanda di Apertura della Liquidazione Giudiziale.docx',
    '099 - Intimazione di sfratto per finita locazione e contestuale citazione per convalida.docx',
    'Conclusioni scritte della parte civile.docx',
    'Delega a sostituto processuale.docx',
    'Domanda Riparazione Ingiusta Detenzione.docx',
    'Lettera di comunicazione di nomina quale difensore d\'ufficio.docx',
    'Memoria Difensiva ex Art. 415-bis c.p.p_.docx',
    'Nomina Difensore Parte Civile.docx',
    'Nota Spese Parte Civile.docx',
    'Nuovi Motivi Ricorso Cassazione.docx',
    'Opposizione a decreto penale di condanna.docx',
    'Procura speciale per costituzione di parte civile.docx',
    'Querela per Atti Persecutori (Stalking).docx',
    'Querela per violenza privata.docx',
    'Richiesta di Affidamento in Prova al Servizio Sociale.docx',
    'Richiesta di Continuazione tra Reati.docx',
    'Ricorso per Cassazione - Processo Penale.docx',
]

with open(CATALOG_PATH) as f:
    catalog = json.load(f)

existing_filenames = {e['filename'] for e in catalog}
added = []
failed = []

for filename in unmatched:
    new_path = NEW_DIR / filename
    if not new_path.exists():
        log(f"SKIP (file not found): {filename}")
        continue

    title = clean_title(filename)
    log(f"\nProcessing: {title}")

    try:
        text = extract_text_from_docx(new_path)
        metadata = get_metadata_from_llm(title, text)

        codice = metadata['codice']
        prefix = codice_prefix(codice)
        tipo_slug = slugify(metadata['tipo_atto'])
        target_filename = f"{prefix}__{tipo_slug}.docx"

        # Avoid filename collisions
        counter = 1
        base_filename = target_filename
        while target_filename in existing_filenames:
            target_filename = f"{prefix}__{tipo_slug}-{counter}.docx"
            counter += 1

        sections = extract_sections_from_docx(new_path)

        entry = {
            'filename': target_filename,
            'codice': codice,
            'categorie': metadata.get('categorie', []),
            'sottocategorie': metadata.get('sottocategorie', []),
            'tipo_atto': metadata['tipo_atto'],
            'label': metadata.get('label', metadata['tipo_atto']),
            'description': metadata.get('description', ''),
            'fields': metadata.get('fields', []),
            'flagged_issues': [],
            'sections': sections,
        }

        # Copy DOCX to templates dir
        target_path = TEMPLATES_DIR / target_filename
        shutil.copy2(new_path, target_path)

        catalog.append(entry)
        existing_filenames.add(target_filename)
        added.append((filename, target_filename, codice))
        log(f"  OK: → {target_filename} | {codice} | {len(sections)} sections")

    except Exception as e:
        log(f"  FAILED: {e}")
        failed.append(filename)

# Save catalog
with open(CATALOG_PATH, 'w', encoding='utf-8') as f:
    json.dump(catalog, f, ensure_ascii=False, indent=2)

log(f"\n{'='*60}")
log(f"Added: {len(added)}, Failed: {len(failed)}")
if failed:
    log("Failed files:")
    for f in failed:
        log(f"  {f}")
