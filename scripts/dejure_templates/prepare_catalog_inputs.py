"""Step 1 of 3 - everything that needs no model, on this PC.

For each fixed template (converted_all_final) decide what goes in the catalog and
build the non-LLM part of its entry: label, area (codice), DEJURE collection,
filename, sections from the document, flags. Applies the decisions taken:
  - skip the 102 titles already live from the pilot
  - mislabelled downloads: rename to the form they contain, or drop (mismatch_plan.csv)
  - 12 titles that exist among the original 484: 2 replace the old entry,
    10 are kept alongside it with their area in the label
Writes enrich_inputs.json (+ categories per area), read by enrich_run.py, which
only makes the model calls and can run on the server.
"""
import collections
import csv
import json
import multiprocessing as mp
import os
import re
import sys
import unicodedata as ud

BASE = r"C:\Users\anton\Downloads\downloads\downloads"
FINAL = os.path.join(BASE, "converted_all_final")
PILOT = os.path.join(BASE, "converted")
HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.environ.get("DEJURE_DATA", r"C:\Users\anton\Downloads\catalog_enrich")   # working files, outside the repo
CATALOG = os.path.join(DATA, "catalog_production.json")
OUT = os.path.join(DATA, "enrich_inputs.json")

sys.path.insert(0, HERE)
from enrich_dejure_templates import (  # noqa: E402
    COLLECTION_TO_CODICE, CODICE_PREFIX, SECTIONS_BUDGET, LONG_TEMPLATE,
    act_text, is_heading, slugify, trim_sections,
)

# Court collections default to civil procedure; their contracts, letters and
# clauses go to the out-of-court area. (The pilot mapped AGRARIO and CONDOMINIO
# wholesale to "Contratti", but most of their ~350 forms are ricorsi/opposizioni.)
COLLECTION_TO_CODICE = dict(COLLECTION_TO_CODICE)
COLLECTION_TO_CODICE["FORMULARIO AGRARIO"] = "Codice di procedura civile"
COLLECTION_TO_CODICE["FORMULARIO CONDOMINIO E LOCAZIONE"] = "Codice di procedura civile"
COURT_COLLECTIONS = {"FORMULARIO DEL CONTENZIOSO BANCARIO", "FORMULARIO AGRARIO",
                     "FORMULARIO CONDOMINIO E LOCAZIONE"}
OUT_OF_COURT = re.compile(
    r"^(contratto|clausola|accordo|lettera|comunicazione|invito|proposta|"
    r"dichiarazione|avviso|comodato|anticipazione|estratto|preventivo|polizza|"
    r"convenzione|verbale|disdetta|diffida|raccomandata|notifica del locatore|"
    r"accettazione|rifiuto|richiesta|risposta|recesso|designazione)\b", re.I)
CODE_HEADER = {"CODICE DI PROCEDURA CIVILE": "Codice di procedura civile",
               "CODICE CIVILE": "Codice di procedura civile",
               "CODICE DI PROCEDURA PENALE": "Codice di procedura penale"}

# Titles that also exist among the original 484 (user decision, 29 Sept 2026)
KEEP_BOTH = {
    "controricorso": "Controricorso (contenzioso tributario)",
    "istanza di assegnazione": "Istanza di assegnazione (esecuzione immobiliare)",
    "note difensive": "Note difensive (processo del lavoro)",
    "opposizione a decreto ingiuntivo": "Opposizione a decreto ingiuntivo (Sezione specializzata agraria)",
    "opposizione agli atti esecutivi": "Opposizione agli atti esecutivi (esecuzione per rilascio)",
    "procura speciale": "Procura speciale (contenzioso tributario)",
    "ricorso in appello": "Ricorso in appello (Sezione specializzata agraria)",
    "ricorso per cassazione": "Ricorso per cassazione (Sezione specializzata agraria)",
    "ricorso per decreto ingiuntivo": "Ricorso per decreto ingiuntivo (Sezione specializzata agraria)",
    "ricorso per revocazione": "Ricorso per revocazione (Sezione specializzata agraria)",
}
REPLACE_OLD = {"lettera di licenziamento per giusta causa", "istanza di conversione del pignoramento"}


def key(s):
    s = ud.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", re.sub(r"\.(docx|pdf)$", "", s)).strip()


def pdf_info(pdf_path):
    """(collection, code_header) from page 1 of the source PDF."""
    import fitz
    if not os.path.exists(pdf_path):
        return "", ""
    with fitz.open(pdf_path) as d:
        lines = [l.strip() for l in d[0].get_text().splitlines() if l.strip()]
    coll = next((" ".join(l.split()).upper() for l in lines if re.match(r"FORMULARIO\b", l, re.I)), "")
    code = next((CODE_HEADER[l.upper()] for l in lines[:3] if l.upper() in CODE_HEADER), "")
    return coll, code


def sections_of(docx_path):
    """Heading + text under it, from the template; note paragraphs left out."""
    from docx import Document
    from convert_dejure_templates import find_note_blocks
    doc = Document(docx_path)
    paras = [p for p in doc.paragraphs if p.text.strip()]
    _, blocks = find_note_blocks(paras)
    notes = {i for b in blocks for i in b}
    sections, heading, buf = [], "Intestazione", []
    for i, par in enumerate(paras):
        if i in notes:
            continue
        if is_heading(par):
            if buf:
                sections.append({"heading": heading, "content": "\n".join(buf)})
            heading, buf = par.text.strip(), []
        else:
            buf.append(par.text.strip())
    if buf:
        sections.append({"heading": heading, "content": "\n".join(buf)})
    # tables: python-docx paragraphs skip them - add their text as a section
    for t in doc.tables:
        rows = [" | ".join(" ".join(c.text.split()) for c in r.cells) for r in t.rows]
        rows = [r for r in rows if r.strip(" |")]
        if rows:
            sections.append({"heading": "Tabella", "content": "\n".join(rows)})
    return [s for s in sections if s["content"].strip()]


def build(job):
    docx_name, pdf_name, label = job
    coll, code = pdf_info(os.path.join(BASE, os.path.splitext(pdf_name)[0] + ".pdf"))
    if coll:
        codice = COLLECTION_TO_CODICE.get(coll, "Contratti e Atti Stragiudiziali")
        if coll in COURT_COLLECTIONS and OUT_OF_COURT.match(label):
            codice = "Contratti e Atti Stragiudiziali"
    else:
        codice = code or "Codice di procedura civile"
    sections, trimmed = trim_sections(sections_of(os.path.join(FINAL, docx_name)))
    flagged = []
    if trimmed:
        flagged.append("sezioni abbreviate (modello oltre %d caratteri)" % SECTIONS_BUDGET)
    elif sum(len(s["content"]) for s in sections) > LONG_TEMPLATE:
        flagged.append("modello lungo: sezioni complete nel prompt di generazione")
    if len(act_text(sections)) < 400:
        flagged.append("testo molto breve")
    return {"source_docx": docx_name, "label": label, "codice": codice, "collection": coll or code,
            "sections": sections, "flagged_issues": flagged}


def main():
    catalog = json.load(open(CATALOG, encoding="utf-8"))
    fixlog = list(csv.DictReader(open(os.path.join(DATA, "fix_converted_all.csv"), encoding="utf-8-sig")))
    pdf_for = {r["new_name"]: r["file"] for r in fixlog}
    plan = {r["docx"]: r for r in csv.DictReader(open(os.path.join(DATA, "mismatch_plan.csv"), encoding="utf-8-sig"))}
    pilot = {key(n) for n in os.listdir(PILOT) if n.endswith(".docx")}
    old_by_key = {key(e.get("label", "")): e for e in catalog[:484]}

    jobs, skipped = [], collections.Counter()
    for name in sorted(n for n in os.listdir(FINAL) if n.endswith(".docx")):
        title = os.path.splitext(name)[0]
        p = plan.get(name)
        if p and p["action"] == "drop":
            skipped["mislabelled copy, dropped"] += 1
            continue
        if p and p["action"] == "rename":
            title = os.path.splitext(p["rename_to"])[0]
        k = key(title)
        if k in pilot:
            skipped["already live from the pilot"] += 1
            continue
        if k in KEEP_BOTH:
            title = KEEP_BOTH[k]
        jobs.append((name, pdf_for.get(name, name), title))

    with mp.Pool(max(1, (os.cpu_count() or 2) - 1)) as pool:
        entries = pool.map(build, jobs, chunksize=8)

    existing = {e["filename"] for e in catalog}
    used = set(existing)
    for i, e in enumerate(entries):
        fn = "%s__%s.docx" % (CODICE_PREFIX[e["codice"]], slugify(e["label"]))
        base, n = fn[:-5], 2
        while fn in used:
            fn = "%s-%d.docx" % (base, n)
            n += 1
        used.add(fn)
        e["id"] = i
        e["filename"] = fn
        k = key(e["label"])
        e["replaces"] = old_by_key[k]["filename"] if k in REPLACE_OLD and k in old_by_key else ""

    cats = collections.defaultdict(set)
    for e in catalog:
        cats[e.get("codice", "")].update(c for c in e.get("categorie", []) if c)
    json.dump({"entries": entries, "categories": {k: sorted(v) for k, v in cats.items()}},
              open(OUT, "w", encoding="utf-8"), ensure_ascii=False)

    print("templates to add: %d  (skipped: %s)" % (len(entries), dict(skipped)))
    print("by area:", dict(collections.Counter(e["codice"] for e in entries).most_common()))
    print("replacing old entries:", [e["replaces"] for e in entries if e["replaces"]])
    print("labels with the area added:", sum(1 for e in entries if e["label"] in KEEP_BOTH.values()))
    lens = sorted(len(e["filename"]) for e in entries)
    print("filename length: max %d, over 200: %d" % (lens[-1], sum(1 for x in lens if x > 200)))
    print("flags:", dict(collections.Counter(f.split(" (")[0].split(":")[0] for e in entries for f in e["flagged_issues"])))
    print("-> %s (%.1f MB)" % (OUT, os.path.getsize(OUT) / 1e6))


if __name__ == "__main__":
    main()
