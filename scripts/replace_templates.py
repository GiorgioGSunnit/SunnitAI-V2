import json
import shutil
import re
from pathlib import Path
from datetime import datetime
from rapidfuzz import fuzz, process
from docx import Document

CATALOG_PATH = Path('/opt/chatbot/data/system_templates/catalog_enriched.json')
TEMPLATES_DIR = Path('/opt/chatbot/data/system_templates')
NEW_TEMPLATES_DIR = Path('/opt/chatbot/data/new_templates')
BACKUP_DIR = Path('/opt/chatbot/data/system_templates_backup')
LOG_PATH = Path('/tmp/replace_templates.log')

BACKUP_DIR.mkdir(parents=True, exist_ok=True)

def log(msg):
    ts = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    line = f"[{ts}] {msg}"
    print(line)
    with open(LOG_PATH, 'a') as f:
        f.write(line + '\n')

def extract_sections_from_docx(docx_path):
    """Extract sections from a DOCX file as heading/content pairs."""
    doc = Document(docx_path)
    sections = []
    current_heading = None
    current_content = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        # Treat bold or ALL CAPS short lines as headings
        is_heading = (
            para.style.name.startswith('Heading') or
            (len(text) < 80 and text.isupper()) or
            (para.runs and all(r.bold for r in para.runs if r.text.strip()))
        )
        if is_heading:
            if current_heading and current_content:
                sections.append({
                    'heading': current_heading,
                    'content': '\n'.join(current_content).strip()
                })
            current_heading = text
            current_content = []
        else:
            current_content.append(text)

    if current_heading and current_content:
        sections.append({
            'heading': current_heading,
            'content': '\n'.join(current_content).strip()
        })

    # If no headings detected, treat whole doc as one section
    if not sections:
        all_text = '\n'.join(p.text.strip() for p in doc.paragraphs if p.text.strip())
        sections = [{'heading': 'Contenuto', 'content': all_text}]

    return sections

def clean_title(filename):
    """Extract clean title from filename like '028 - Atto di appello.docx'"""
    name = Path(filename).stem
    # Remove leading number and dash
    name = re.sub(r'^\d+\s*[-–]\s*', '', name)
    return name.strip()

def main():
    with open(CATALOG_PATH) as f:
        catalog = json.load(f)

    catalog_labels = [e.get('label') or e.get('tipo_atto', '') for e in catalog]
    new_files = sorted(NEW_TEMPLATES_DIR.glob('*.docx'))

    log(f"Found {len(new_files)} new DOCX files")
    log(f"Catalog has {len(catalog)} entries")
    log("=" * 60)

    replaced = []
    no_match = []

    for new_file in new_files:
        title = clean_title(new_file.name)
        
        # Fuzzy match against catalog labels
        match = process.extractOne(
            title,
            catalog_labels,
            scorer=fuzz.token_sort_ratio,
            score_cutoff=70
        )

        if not match:
            log(f"NO MATCH: '{title}' (score too low)")
            no_match.append(new_file.name)
            continue

        matched_label, score, idx = match
        entry = catalog[idx]
        old_filename = entry['filename']
        old_path = TEMPLATES_DIR / old_filename

        log(f"MATCH: '{title}' → '{matched_label}' (score={score:.0f})")
        log(f"  Catalog filename: {old_filename}")

        # Backup old DOCX
        if old_path.exists():
            backup_path = BACKUP_DIR / old_filename
            shutil.copy2(old_path, backup_path)
            log(f"  Backed up: {old_path} → {backup_path}")

        # Copy new DOCX in place
        shutil.copy2(new_file, old_path)
        log(f"  Replaced: {new_file.name} → {old_path}")

        # Extract sections from new DOCX
        new_sections = extract_sections_from_docx(new_file)
        old_section_count = len(entry.get('sections', []))
        entry['sections'] = new_sections
        log(f"  Sections: {old_section_count} → {len(new_sections)}")

        replaced.append((new_file.name, matched_label, score))

    # Write updated catalog
    with open(CATALOG_PATH, 'w', encoding='utf-8') as f:
        json.dump(catalog, f, ensure_ascii=False, indent=2)
    log("=" * 60)
    log(f"Catalog updated: {CATALOG_PATH}")
    log(f"Replaced: {len(replaced)}/{len(new_files)}")
    log(f"No match: {len(no_match)}")
    if no_match:
        log("Files with no catalog match (new types — handle separately):")
        for f in no_match:
            log(f"  {f}")

if __name__ == '__main__':
    main()
