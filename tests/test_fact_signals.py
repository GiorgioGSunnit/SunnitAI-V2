"""FACT_SIGNALS_ENABLED (src/rag/fact_signals.py, fact_signal_rules.yaml): date arithmetic, typed facts,
one invented case per rule. Synthetic facts only -- never the held-out texts."""

import datetime as dt
import json
import os

import pytest

os.environ.setdefault("NEO4J_URI", "bolt://localhost:7687")
os.environ.setdefault("NEO4J_USER", "neo4j")
os.environ.setdefault("NEO4J_PASSWORD", "test-password")
os.environ.setdefault("LLM_API_KEY", "test-key")
os.environ.setdefault("EMBEDDING_PROVIDER", "openai")
os.environ.setdefault("EMBEDDING_API_KEY", "test-key")

import src.rag.defensive_generation as dg  # noqa: E402
import src.rag.plan_then_write as ptw  # noqa: E402
from src.rag import fact_signals as fs  # noqa: E402

TODAY = "2026-10-27"


@pytest.fixture(autouse=True)
def fixed_today(monkeypatch):
    monkeypatch.setenv(fs.TODAY_ENV, TODAY)
    monkeypatch.delenv(fs.FACT_SIGNALS_FLAG, raising=False)
    monkeypatch.delenv(fs.REGIMES_FLAG, raising=False)


def F(**kw):
    base = {k: ([] if isinstance(v, list) else "") for k, v in fs.SCHEMA.items()}
    base.update(kw)
    return base


def ev(what, date=""):
    return {"date": date, "what": what, "who": ""}


def party(name, role, age=""):
    return {"name": name, "role": role, "age": age}


def contract(type_, clauses=""):
    return {"type": type_, "parties": "", "clauses": clauses}


def proc(stage, decision="", date="", authority=""):
    return {"stage": stage, "authority": authority, "decision": decision, "date": date}


# --- dates and terms -------------------------------------------------------------------------------

def test_dates_as_written_with_their_precision():
    assert fs.parse_date("il 12/03/2016") == (dt.date(2016, 3, 12), dt.date(2016, 3, 12), "2016-03-12")
    assert fs.parse_date("1° marzo 2016")[2] == "2016-03-01"
    assert fs.parse_date("nel marzo del 2016") == (dt.date(2016, 3, 1), dt.date(2016, 3, 31), "2016-03")
    assert fs.parse_date("nel 2016") == (dt.date(2016, 1, 1), dt.date(2016, 12, 31), "2016")
    assert fs.parse_date("2016-03-12")[2] == "2016-03-12"
    assert fs.parse_date("tre mesi dopo") is None and fs.parse_date("31/02/2016")[2] == "2016"   # no such day: its year


def test_periods_and_status_respect_the_precision():
    assert fs.add_period(dt.date(2024, 2, 29), years=1) == dt.date(2025, 2, 28)
    assert fs.add_period(dt.date(2026, 1, 31), months=1) == dt.date(2026, 2, 28)
    year_2016, year_2018 = fs.parse_date("2016")[:2], fs.parse_date("2018")[:2]
    assert fs.term_status(year_2016, year_2018, years=1) == "DECORSO"         # even 31 Dec 2016 -> 1 Jan 2018
    assert fs.term_status(year_2016, fs.parse_date("2017")[:2], years=1) == "INCERTO"
    day = fs.parse_date("1 marzo 2026")[:2]
    assert fs.term_status(day, fs.parse_date("20 marzo 2026")[:2], months=3) == "NON DECORSO"


def test_the_possessory_term_reads_as_the_lawyer_writes_it():
    terms = fs.compute_terms(F(events=[ev("Caio chiude il cancello del passaggio", "2016"),
                                       ev("Tizio propone ricorso possessorio", "2018")]))
    t = next(t for t in terms if t["id"] == "possessoria")
    assert t["status"] == "DECORSO" and t["ref"] == "artt. 1168 e 1170 c.c."
    assert t["text"] == ("azione possessoria (spoglio o turbativa): termine annuale DECORSO (Caio chiude il "
                         "cancello del passaggio 2016, Tizio propone ricorso possessorio 2018)")


def test_a_term_with_no_act_runs_to_today():
    terms = fs.compute_terms(F(events=[ev("Caio occupa il cortile", "1 settembre 2026")]))
    t = next(t for t in terms if t["id"] == "possessoria")
    assert t["status"] == "NON DECORSO" and t["to"] == {"what": "oggi", "date": TODAY}


def test_criminal_appeal_terms_follow_the_deposit_of_the_reasons():
    sentence = proc("primo grado", "sentenza di condanna", "10 settembre 2026")
    t = next(t for t in fs.compute_terms(F(proceedings=[sentence])) if t["id"] == "impugnazione_penale")
    assert t["status"] == "INCERTO" and "dipende da quando è depositata la motivazione" in t["text"]
    deposited = F(proceedings=[sentence], events=[ev("deposito della motivazione", "1 ottobre 2026")])
    assert next(t for t in fs.compute_terms(deposited) if t["id"] == "impugnazione_penale")["status"] == "NON DECORSO"
    late = F(proceedings=[proc("primo grado", "sentenza di condanna", "10 gennaio 2026")])
    assert next(t for t in fs.compute_terms(late) if t["id"] == "impugnazione_penale")["status"] == "DECORSO"


def test_usucapione_and_tender_terms():
    held = F(possession=[{"who": "Tizio", "asset": "fondo di Caio", "since": "1995", "how": "pacifico",
                          "still_ongoing": "sì"}])
    assert next(t for t in fs.compute_terms(held) if t["id"] == "usucapione_ordinaria")["status"] == "DECORSO"
    car = F(possession=[{"who": "Tizio", "asset": "auto di Caio", "since": "2020", "how": "", "still_ongoing": "sì"}])
    assert [t["id"] for t in fs.compute_terms(car)] == ["usucapione_mobili"]
    tender = F(proceedings=[proc("gara", "esclusione dalla gara", "1 ottobre 2026")])
    ids = {t["id"]: t["status"] for t in fs.compute_terms(tender)}
    assert ids == {"ricorso_tar_appalti": "NON DECORSO"}


# --- typed facts ---------------------------------------------------------------------------------

def test_a_date_the_text_does_not_hold_is_dropped():
    case = "Nel 2016 Caio chiuse il cancello. Tizio agì due anni dopo."
    raw = {"events": [{"date": "2016", "what": "Caio chiude il cancello"}, {"date": "2018", "what": "Tizio agisce"}],
           "parties": [{"name": "Tizio", "role": "cliente", "age": 40}], "extra": "dropped", "requests": "Che fare?"}
    facts = fs.validate_facts(raw, case)
    assert facts["events"][0]["date"] == "2016" and facts["events"][0]["date_iso"] == "2016"
    assert facts["events"][1]["date"] == "" and facts["events"][1]["date_iso"] == ""      # 2018 is not in the text
    assert facts["parties"][0]["age"] == "40" and facts["requests"] == ["Che fare?"] and "extra" not in facts
    assert fs.validate_facts(["not", "a", "dict"], case) is None and fs.validate_facts({"foo": 1}, case) is None


def test_extraction_is_one_call_retried_once(monkeypatch):
    replies = iter(["non è JSON", json.dumps({"events": [{"date": "", "what": "x"}]})])
    calls = []

    def chat(messages, max_tokens=None, temperature=None, **kw):
        calls.append((messages[0].content, temperature))
        return next(replies)
    monkeypatch.setattr(fs, "_call_chat", chat)
    facts, problem = fs.extract_facts("Un caso.")
    assert facts["events"][0]["what"] == "x" and problem == "" and len(calls) == 2
    assert calls[0][1] == 0 and calls[1][0].endswith(fs._RETRY_SUFFIX)


# --- the rules: one invented case each -------------------------------------------------------------

FIXTURES = {
    "destinazione_padre_famiglia": F(property_history=[{"asset": "fondo A", "owner": "Mevio"}, {"asset": "fondo B", "owner": "Mevio"}],
                                     works_or_signs=[{"what": "strada sterrata", "visible": "sì", "existing_since": ""}]),
    "usucapione_servitu_apparente": F(works_or_signs=[{"what": "sentiero", "visible": "sì", "existing_since": "1990"}],
                                      possession=[{"who": "Tizio", "asset": "sentiero", "since": "1990", "how": "pacifico",
                                                   "still_ongoing": "sì"}]),
    "usucapione_ordinaria": F(possession=[{"who": "Tizio", "asset": "orto", "since": "1995", "how": "pacifico e pubblico",
                                           "still_ongoing": "sì"}]),
    "possesso_vizioso": F(possession=[{"who": "Tizio", "asset": "orto", "since": "2010", "how": "clandestino",
                                       "still_ongoing": "sì"}]),
    "possessoria_preclusa": F(events=[ev("Caio chiude il cancello del passaggio", "marzo 2016"),
                                      ev("Tizio propone ricorso possessorio", "2018")]),
    "possessoria_in_termine": F(events=[ev("Caio chiude il cancello", "1 settembre 2026")]),
    "fondo_intercluso": F(events=[ev("il fondo di Tizio è intercluso")]),
    "immissioni": F(events=[ev("odori e fumi dal ristorante vicino")]),
    "distanze_vedute": F(events=[ev("Caio costruisce un muro a un metro dal confine")]),
    "condominio_parti_comuni": F(events=[ev("l'assemblea del condominio vieta l'uso del cortile")]),
    "clausola_penale": F(contracts=[contract("appalto", "penale per ritardo")], events=[ev("ritardo di tre mesi")]),
    "caparra": F(contracts=[contract("preliminare di vendita", "caparra confirmatoria")]),
    "risoluzione_contratto": F(contracts=[contract("vendita")], requests=["Può risolvere il contratto?"]),
    "vizi_vendita": F(contracts=[contract("compravendita")], events=[ev("la caldaia ha difetti")]),
    "conformita_consumatore": F(parties=[party("Tizio", "consumatore"), party("Alfa srl", "impresa")],
                                contracts=[contract("compravendita di uno smartphone")],
                                events=[ev("lo smartphone non funziona")]),
    "clausole_vessatorie": F(parties=[party("Tizio", "consumatore")],
                             contracts=[contract("abbonamento", "foro esclusivo dell'impresa")]),
    "usura": F(money=[{"amount": "10000", "rate": "25", "period": "annuo", "threshold": "16"}]),
    "contatto_sociale_scuola": F(parties=[party("Marco", "minore", "8")],
                                 events=[ev("durante la ricreazione a scuola Marco cade")],
                                 injuries=[{"who": "Marco", "days_to_heal": "20", "permanent": "no"}]),
    "responsabilita_sanitaria_civile": F(parties=[party("dott. Rossi", "medico")],
                                         events=[ev("intervento chirurgico al ginocchio")],
                                         injuries=[{"who": "Tizio", "days_to_heal": "90", "permanent": "sì"}]),
    "responsabilita_sanitaria_penale": F(parties=[party("dott. Rossi", "medico")],
                                         injuries=[{"who": "Tizio", "days_to_heal": "90", "permanent": "sì"}],
                                         proceedings=[proc("indagini preliminari", authority="Procura")]),
    "albergatore": F(events=[ev("dalla camera d'albergo sono stati rubati i gioielli")]),
    "difformita_appalto": F(contracts=[contract("appalto di ristrutturazione")], events=[ev("infiltrazioni nel bagno")]),
    "rovina_gravi_difetti": F(contracts=[contract("appalto")], events=[ev("lavori ultimati", "2015"),
                                                                       ev("compaiono crepe nei muri portanti", "2025")]),
    "circolazione_veicoli": F(events=[ev("incidente stradale all'incrocio")]),
    "prescrizione_decorsa": F(events=[ev("incidente stradale", "2015")]),
    "tentativo": F(events=[ev("Tizio si è procurato un piede di porco e ha fatto un sopralluogo")]),
    "concorso_persone": F(parties=[party("Tizio", "indagato/imputato"), party("Caio", "indagato/imputato")],
                          proceedings=[proc("indagini", authority="Procura")]),
    "concorso_extraneus": F(parties=[party("Rossi", "pubblico ufficiale, indagato/imputato"),
                                     party("Bianchi", "indagato/imputato")], proceedings=[proc("indagini")]),
    "morte_o_lesioni_da_altro_delitto": F(events=[ev("Tizio spinge Caio durante una lite")],
                                          injuries=[{"who": "Caio", "days_to_heal": "", "permanent": "morte"}]),
    "medico_e_latitante": F(parties=[party("dott. Verdi", "medico")],
                            events=[ev("il medico cura un uomo ricercato, ferito da colpo di arma da fuoco"),
                                    ev("il medico non ha avvisato le autorità")]),
    "querela_tardiva": F(events=[ev("Caio minaccia Tizio", "gennaio 2026")]),
    "impugnazione_penale": F(proceedings=[proc("primo grado", "sentenza di condanna", "10 settembre 2026")]),
    "legittima_difesa": F(events=[ev("Tizio, aggredito da Caio, ha reagito colpendolo")]),
    "ottemperanza": F(proceedings=[proc("ricorso", "sentenza di annullamento del diniego", authority="TAR")],
                      events=[ev("il Comune non ha dato esecuzione alla sentenza")]),
    "riadozione_senza_contraddittorio": F(events=[ev("il provvedimento annullato è stato riadottato con motivazione "
                                                     "per relationem")]),
    "esclusione_gara": F(events=[ev("la società è stata esclusa dalla gara")]),
    "ricorso_tar_decorso": F(parties=[party("Comune di X", "pubblica amministrazione")],
                             proceedings=[proc("provvedimento", "diniego del permesso", "1 marzo 2026")]),
    "silenzio_amministrazione": F(parties=[party("Comune di X", "pubblica amministrazione")],
                                  events=[ev("il Comune non ha risposto all'istanza")]),
    "comunione_de_residuo": F(events=[ev("durante il matrimonio la moglie acquista un locale per la sua impresa")]),
    "mantenimento": F(events=[ev("separazione dei coniugi; il padre non versa il mantenimento")]),
    "accettazione_tacita": F(parties=[party("Anna", "erede")],
                             events=[ev("Anna riscuote i canoni dell'appartamento del defunto")]),
    "lesione_legittima": F(events=[ev("con testamento il padre lascia tutto al figlio maggiore, escludendo la figlia")]),
    "morosita_conduttore": F(contracts=[contract("locazione abitativa")], events=[ev("il conduttore non paga i canoni")]),
    "finita_locazione": F(contracts=[contract("locazione")],
                          events=[ev("alla scadenza il conduttore non vuole lasciare l'immobile")]),
    "licenziamento": F(parties=[party("Alfa srl", "datore di lavoro"), party("Tizio", "lavoratore")],
                       events=[ev("il lavoratore viene licenziato")]),
    "differenze_retributive": F(parties=[party("Alfa srl", "datore di lavoro"), party("Tizio", "lavoratore")],
                                events=[ev("gli straordinari non sono stati pagati")]),
}


def _fired(facts):
    return {r["id"] for r in fs.fire_rules(facts, fs.compute_terms(facts))}


def test_every_rule_has_an_invented_case():
    assert {r["id"] for r in fs.load_rules()} == set(FIXTURES)


@pytest.mark.parametrize("rule_id", sorted(FIXTURES))
def test_each_rule_fires_on_its_case(rule_id):
    assert rule_id in _fired(FIXTURES[rule_id])


def test_rules_that_must_not_fire():
    violent = F(possession=[{"who": "Tizio", "asset": "orto", "since": "1995", "how": "violento", "still_ongoing": "sì"}])
    assert "usucapione_ordinaria" not in _fired(violent) and "possesso_vizioso" in _fired(violent)
    recent = F(possession=[{"who": "Tizio", "asset": "orto", "since": "2015", "how": "pacifico", "still_ongoing": "sì"}])
    assert "usucapione_ordinaria" not in _fired(recent)
    assert "possessoria_in_termine" not in _fired(FIXTURES["possessoria_preclusa"])
    assert "responsabilita_sanitaria_civile" not in _fired(FIXTURES["responsabilita_sanitaria_penale"])
    assert "condominio_parti_comuni" not in _fired(F(events=[ev("l'assemblea dei soci approva il bilancio")]))
    assert "ottemperanza" not in _fired(F(proceedings=[proc("primo grado", "sentenza di condanna", authority="Tribunale")],
                                          events=[ev("non ha eseguito la sentenza")]))
    criminal = {r["area"] for r in fs.fire_rules(FIXTURES["concorso_persone"], [])}
    assert criminal == {"penale"}
    assert _fired(F()) == set() and fs.fire_rules(None, []) == []


def test_rules_name_institutions_never_articles(tmp_path):
    for rule in fs.load_rules():
        assert rule["status"] == "experiment" and rule["codice"]
    bad = tmp_path / "bad.yaml"
    bad.write_text("- {id: x, area: penale, when: [{fact: {pattern: 'a'}}], institution: 'art. 56 c.p.', "
                   "codice: c.p., why: y, status: experiment}\n", encoding="utf-8")
    with pytest.raises(ValueError, match="names an article"):
        fs.load_rules(str(bad))
    bad.write_text("- {id: x, area: penale, when: [{guess: 1}], institution: tentativo, codice: c.p., why: y, "
                   "status: experiment}\n", encoding="utf-8")
    with pytest.raises(ValueError, match="unknown condition"):
        fs.load_rules(str(bad))


def test_central_issues_once_each_and_capped():
    fired = fs.fire_rules(FIXTURES["possessoria_preclusa"], fs.compute_terms(FIXTURES["possessoria_preclusa"]))
    central = fs.central_issues(fired + fired)
    assert central[0] == {"nome": "decadenza dall'azione di reintegrazione e di manutenzione", "codice": "c.c.",
                          "fonte": "regola possessoria_preclusa", "ranges": [[1140, 1172], [948, 951], [1079, 1079]]}
    assert len(central) == len({c["nome"] for c in central}) <= fs.MAX_CENTRAL


# --- integration ---------------------------------------------------------------------------------

def test_fired_institutions_join_the_frames_central_issues(monkeypatch):
    monkeypatch.setenv("CASE_FRAME_ENABLED", "true")
    monkeypatch.setenv("FACT_SIGNALS_ENABLED", "true")
    monkeypatch.setattr(dg, "_case_frame", lambda facts, request: {
        "tasks": ["Che cosa può fare Tizio?"], "central_issue_candidates": [{"nome": "usucapione", "codice": "c.c."}]})
    monkeypatch.setattr(fs, "extract_facts", lambda text: (FIXTURES["usucapione_ordinaria"], ""))
    frame, got = dg._frame_and_fact_signals("Tizio coltiva l'orto di Caio dal 1995.", "")
    assert [r["id"] for r in got["fired"]] == ["usucapione_ordinaria"]
    assert [c["nome"] for c in frame["central_issue_candidates"]] == ["usucapione"]      # already there: once
    monkeypatch.setattr(fs, "extract_facts", lambda text: (FIXTURES["fondo_intercluso"], ""))
    frame, _ = dg._frame_and_fact_signals("x", "")
    assert [c["nome"] for c in frame["central_issue_candidates"]] == [
        "usucapione", "passaggio coattivo", "indennità per il passaggio coattivo"]
    assert frame["central_issue_candidates"][1]["ranges"] == [[1051, 1055]]
    analysis = dg._with_central_questions({"questioni": []}, frame)
    assert analysis["questioni"][1]["istituti"][0] == {"nome": "passaggio coattivo", "codice": "c.c.", "numero_indicativo": "",
                                                       "fonte": "regola fondo_intercluso", "ranges": [[1051, 1055]]}


def test_rules_fire_only_within_the_case_areas(monkeypatch):
    monkeypatch.setenv("CASE_FRAME_ENABLED", "true")
    facts = FIXTURES["legittima_difesa"]
    assert "legittima_difesa" in {r["id"] for r in fs.fire_rules(facts, [], areas=["penale"])}
    assert fs.fire_rules(facts, [], areas=["contratti", "responsabilita_civile"]) == []
    assert "legittima_difesa" in {r["id"] for r in fs.fire_rules(facts, [], areas=["altro"])}         # unknown: no gate
    usura = FIXTURES["usura"]
    assert {r["id"] for r in fs.fire_rules(usura, [], areas=["penale"])} == {"usura"}                 # both families
    monkeypatch.setattr(dg, "_case_frame", lambda facts, request: {"areas": ["contratti"], "central_issue_candidates": []})
    monkeypatch.setattr(fs, "extract_facts", lambda text: (facts, ""))
    frame, got = dg._frame_and_fact_signals("x", "")
    assert got["fired"] == [] and got["areas"] == ["contratti"]


def test_computed_terms_reach_the_plan_and_the_procedural_section():
    terms = fs.compute_terms(FIXTURES["possessoria_preclusa"])
    analysis = {"fact_signals": {"terms": terms}}
    human = ptw._plan_human("Fatti.", "", analysis, "NORME", "", "", [])
    assert "TERMINI CALCOLATI DALLE DATE DEL CASO:\n- azione possessoria (spoglio o turbativa): termine annuale DECORSO" in human
    assert "non va proposta come praticabile" in human
    plan = {"parte_assistita": "Tizio", "posizione": "", "obiettivo": "", "questione_centrale": "", "risposte": [],
            "questioni": [{"titolo": "Passaggio", "norme": [], "pronunce": [], "argomento": "x", "obiezione": "",
                           "replica": ""}], "subordinate": [], "punti_deboli": []}
    jobs = ptw.section_jobs(plan, "Fatti.", "", analysis, [], {}, "")
    title, instruction, material, _ = jobs[4]
    assert title == "Profili processuali e termini" and "TERMINI CALCOLATI DALLE DATE DEL CASO:\n- azione possessoria" in material
    assert "non va proposta come praticabile" in instruction
    assert "TERMINI CALCOLATI" not in ptw._plan_human("Fatti.", "", {}, "NORME", "", "", [])          # flag off: nothing


def _script(name):
    import importlib.util
    from pathlib import Path
    spec = importlib.util.spec_from_file_location(name, Path(__file__).resolve().parent.parent / "scripts" / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_the_lawyers_table_lists_every_rule_in_italian():
    text = _script("export_fact_rules").markdown(fs)
    assert text.count("| in sperimentazione (approvata per l'esperimento) |") == len(fs.load_rules())
    assert ("| fondo_intercluso | Diritti reali e possesso | un fondo senza accesso alla via pubblica | passaggio "
            "coattivo; indennità per il passaggio coattivo | c.c. |") in text
    assert "| azione possessoria (spoglio o turbativa) | 1 anno | artt. 1168 e 1170 c.c. |" in text


def test_the_bench_counts_lost_articles_a_rule_reaches():
    bench = _script("fact_signals_bench")
    case = {"id": "x1", "slots": [{"label": "art. 1051 cc", "pool_label": "1051 c.c.", "stage": "LOST_AT_UNDERSTANDING"}]}
    facts = FIXTURES["fondo_intercluso"]
    signals = {"facts": facts, "terms": [], "fired": fs.fire_rules(facts, []) + fs.fire_rules(FIXTURES["caparra"], []),
               "problem": ""}
    resolved = {"passaggio coattivo": ["1051 c.c."], "indennità per il passaggio coattivo": ["1053 c.c."],
                "caparra confirmatoria": ["1385 c.c."], "recesso per inadempimento": []}
    result = {"case": "x1", "signals": signals, "resolved": resolved, "lost": case["slots"]}
    ekey = lambda label: (label.split()[0], {"c.c.": "cc"}.get(label.split(" ", 1)[1], label.split(" ", 1)[1]))  # noqa: E731
    text = bench.render("/runs/r", [result], ekey, {"x1": {("1051", "cc")}}, 60, [])
    assert "LOST_AT_UNDERSTANDING in the run: 1; reached by a fired rule: 1 (x1 art. 1051 cc ← fondo_intercluso)" in text
    assert "possible false triggers): 1 (x1 caparra)" in text and "**1051 c.c.**, 1053 c.c." in text
    assert "    - perché: un fondo senza accesso alla via pubblica — «il fondo di tizio è intercluso»" in text


# --- the case's present ------------------------------------------------------------------------------

def test_without_a_present_date_no_term_runs_to_the_real_today(monkeypatch):
    monkeypatch.delenv(fs.TODAY_ENV)
    lonely = F(events=[ev("Caio chiude il cancello", "marzo 2016")])
    t = fs.compute_terms(lonely)[0]
    assert t["status"] == "NON VERIFICABILE" and "il caso non indica la data attuale" in t["text"]
    assert "possessoria_preclusa" not in _fired(lonely) and "possessoria_in_termine" not in _fired(lonely)
    later = F(events=[ev("Caio chiude il cancello", "marzo 2016"), ev("Tizio scrive al vicino", "giugno 2018")])
    t = fs.compute_terms(later)[0]
    assert t["status"] == "DECORSO" and "già al fatto più recente" in t["text"]      # run even by the latest fact
    dated = F(events=[ev("Caio chiude il cancello", "marzo 2016")], present_date="10 aprile 2016")
    assert fs.compute_terms(dated)[0]["status"] == "NON DECORSO"
    assert "usucapione_ordinaria" not in _fired(F(possession=[{"who": "T", "asset": "orto", "since": "2000", "how": "",
                                                                "still_ongoing": "sì"}]))


def test_a_present_date_is_kept_only_as_the_text_writes_it():
    facts = fs.validate_facts({"present_date": "oggi, 10 marzo 2019", "events": []}, "Siamo a oggi, 10 marzo 2019.")
    assert facts["present_date"] == "oggi, 10 marzo 2019"
    assert fs.validate_facts({"present_date": "1 gennaio 2030", "events": []}, "Nessuna data.")["present_date"] == ""


def test_a_cut_reply_keeps_its_complete_values(monkeypatch):
    cut = '{"parties": [{"name": "Tizio", "role": "cliente"}], "events": [{"date": "2016", "what": "a \\"b\\""}, {"date": "20'
    assert fs._parse_or_repair(cut) == {"parties": [{"name": "Tizio", "role": "cliente"}],
                                        "events": [{"date": "2016", "what": 'a "b"'}]}
    assert fs._parse_or_repair('```json\n{"events": []}\n```') == {"events": []}
    monkeypatch.setattr(fs, "_call_chat", lambda messages, **kw: cut)
    facts, problem = fs.extract_facts("Tizio, nel 2016, a.")
    assert facts["parties"][0]["name"] == "Tizio" and problem == ""


# --- why, and which articles ---------------------------------------------------------------------------

def test_a_fired_rule_says_why():
    rule = next(r for r in fs.fire_rules(FIXTURES["fondo_intercluso"], []) if r["id"] == "fondo_intercluso")
    assert rule["matched"] == ["un fondo senza accesso alla via pubblica — «il fondo di tizio è intercluso»"]
    rule = next(r for r in fs.fire_rules(FIXTURES["medico_e_latitante"], []) if r["id"] == "medico_e_latitante")
    assert rule["matched"][0] == "una parte è medico o professionista o struttura sanitaria — «dott. Verdi (medico)»"


def test_a_rule_institution_keeps_only_articles_in_range_whose_rubrica_fits():
    def row(number, rubrica):
        return {"_article": (number, "c.c."), "s": {"titolo_articolo": rubrica}}
    istituto = {"nome": "passaggio coattivo", "codice": "c.c.", "fonte": "regola fondo_intercluso", "ranges": [[1051, 1055]]}
    rows = [row("1051", "Passaggio coattivo"), row("2634", "Infedeltà patrimoniale"), row("1052", "Passaggio coattivo a favore di fondo non intercluso"),
            row("1053", "Indennità")]
    assert [r["_article"][0] for r in fs.keep_resolved(istituto, rows)] == ["1051", "1052"]
    out_of_range = {"nome": "danno non patrimoniale", "fonte": "regola x", "ranges": [[2043, 2059]]}
    assert fs.keep_resolved(out_of_range, [row("2634", "Danno non patrimoniale della società")]) == []
    assert fs.keep_resolved({"nome": "passaggio coattivo"}, rows) == rows                # not a fact rule's: unchanged
    assert fs.keep_resolved({"nome": "accettazione tacita dell'eredità", "fonte": "r"}, [row("476-bis", "Accettazione tacita")])


# --- answers to .docx, blind A/B --------------------------------------------------------------------------

def test_an_answer_and_its_sources_go_to_a_docx(tmp_path):
    from docx import Document
    docx_script = _script("answer_docx")
    record = {"status": "ok", "request": "Che cosa può fare Tizio?",
              "answer": "**1. Risposte ai quesiti**\n\nTizio può agire ex **art. 1051 c.c.**\n- prova del passaggio",
              "citations": [{"document_name": "Codice Civile 2026", "document_type": "primary",
                             "sections": [{"name": "1051.0.0", "title": "Passaggio coattivo", "plain_text": "Il proprietario..."}]},
                            {"document_name": "Cass. civ., sez. II, n. 1/2026", "document_type": "interpretation",
                             "sections": [{"name": "1", "title": "", "plain_text": "Il fondo intercluso ha diritto..."}]}]}
    out = str(tmp_path / "r.docx")
    docx_script.build(record, out, "Risposta 1", record["request"], None)
    text = "\n".join(p.text for p in Document(out).paragraphs)
    for part in ("Risposta 1", "Che cosa può fare Tizio?", "1. Risposte ai quesiti", "Tizio può agire ex art. 1051 c.c.",
                 "prova del passaggio", "Fonti citate", "Norme", "Codice Civile 2026", "Art. 1051 — Passaggio coattivo",
                 "Giurisprudenza", "Il fondo intercluso ha diritto..."):
        assert part in text
    assert "Configurazione" not in text                                          # blind
    docx_script.build(record, out, "Risposta", "", ["FACT_SIGNALS_ENABLED=true"])
    assert "FACT_SIGNALS_ENABLED=true" in "\n".join(p.text for p in Document(out).paragraphs)


def test_the_ab_order_is_random_and_the_settings_parse():
    ab = _script("ab_answers")
    assert ab.settings("A=1, B=2,") == ["A=1", "B=2"] and ab.settings("") == []
    assert ab.assign(1) == ab.assign(1) and sorted(ab.assign(1).values()) == ["A", "B"]
    assert {tuple(ab.assign(seed).values()) for seed in range(20)} == {("A", "B"), ("B", "A")}


# --- FACT_REGIMES_ENABLED (the lawyer's blind review, 26 Oct): invented cases only ---------------------

def person(name, role="", kind="persona fisica", purpose="", occupation="", activities=""):
    return {"name": name, "role": role, "age": "", "kind": kind, "occupation": occupation, "purpose": purpose,
            "activities": activities}


SOFA = F(parties=[person("Anna", "cliente", purpose="personale"), person("Arredi Beta s.r.l.", "controparte", "società")],
         contracts=[contract("compravendita di un divano")],
         events=[ev("dopo la consegna il divano presenta difetti di cucitura")], risks=[])


def test_a_natural_person_buying_for_private_use_from_a_business_is_a_consumer():
    facts, quals = fs.qualify_parties(SOFA)
    assert [(q["party"], q["quality"]) for q in quals] == [("Anna", "consumatore"), ("Arredi Beta s.r.l.", "professionista")]
    assert facts["parties"][0]["role"] == "cliente, consumatore" and facts["_regimes"] is True
    assert "professionista" in facts["parties"][1]["role"]
    # bought for the business, or from another private person: no consumer
    work = F(**{**SOFA, "parties": [person("Anna", purpose="professionale"), SOFA["parties"][1]]})
    assert fs.qualify_parties(work)[1] == []
    private = F(**{**SOFA, "parties": [SOFA["parties"][0], person("Bruno", "venditore")]})
    assert fs.qualify_parties(private)[1] == []
    assert fs.qualify_parties(F(**{**SOFA, "contracts": []}))[1] == []


def test_public_bodies_and_public_officials():
    facts, quals = fs.qualify_parties(F(parties=[
        person("Comune di Alfa", "controparte", "ente pubblico"), person("Carlo", "indagato/imputato",
                                                                          occupation="agente di polizia locale")]))
    assert [(q["party"], q["quality"]) for q in quals] == [("Comune di Alfa", fs.PUBLIC_BODY),
                                                            ("Carlo", fs.PUBLIC_OFFICIAL)]
    assert "pubblico ufficiale" in facts["parties"][1]["role"]
    assert fs.regime_issues(facts, quals, ["penale"])[0]["nome"].startswith("nozione di pubblico ufficiale")
    assert fs.regime_issues(facts, quals, ["responsabilita_civile"]) == []


def test_consumer_regimes_join_the_central_issues_in_civil_cases_only():
    facts, quals = fs.qualify_parties(SOFA)
    names = [r["nome"] for r in fs.regime_issues(facts, quals, ["contratti"])]
    assert names == ["garanzia legale di conformità", "foro del consumatore"]      # no clause: no unfair terms
    assert fs.regime_issues(facts, quals, ["penale"]) == []
    with_clause = F(**{**SOFA, "contracts": [contract("vendita di un divano", "foro esclusivo del venditore")]})
    facts, quals = fs.qualify_parties(with_clause)
    assert fs.regime_issues(facts, quals, [])[-1]["ranges"] == [[33, 38]]


def test_terms_without_dates_listed_only_with_the_flag_and_within_the_case_family(monkeypatch):
    assert fs.compute_terms(SOFA) == []                                   # no dates, no regimes: as before
    facts, _ = fs.qualify_parties(SOFA)
    terms = {t["id"]: t for t in fs.compute_terms(facts, ["contratti"])}
    assert set(terms) == {"denuncia_vizi_vendita", "azione_vizi_vendita", "denuncia_conformita_consumo",
                          "responsabilita_conformita_consumo", "azione_conformita_consumo"}
    assert terms["azione_conformita_consumo"]["text"].startswith(
        "prescrizione dell'azione per il difetto di conformità non dolosamente occultato (consumatore): termine di "
        "26 mesi dalla consegna del bene — DA VERIFICARE")
    assert {t["status"] for t in terms.values()} == {"DA VERIFICARE"}
    offence = F(events=[ev("Tizio danneggia la vetrina del negozio di Caio")], _regimes=True)
    assert [t["id"] for t in fs.compute_terms(offence, ["responsabilita_civile"])] == ["prescrizione_extracontrattuale"]
    assert "querela" in [t["id"] for t in fs.compute_terms(offence, ["penale"])]
    block = fs.terms_block({"terms": list(terms.values())}, refs=True)
    assert "[art. 1495 c.c.]" in block and "[Codice del Consumo]" in block
    assert "[" not in fs.terms_block({"terms": list(terms.values())})


def test_procedural_tools_follow_the_facts():
    facts, _ = fs.qualify_parties(SOFA)
    tools = fs.procedural_tools(facts, ["contratti"])
    assert [t["id"] for t in tools] == ["accertamento_preventivo"]
    assert tools[0]["refs"] == [["696", "c.p.c."], ["696-bis", "c.p.c."]] and "difetti" in tools[0]["matched"][0]
    gas = F(events=[ev("la caldaia installata dalla ditta perde gas")],
            risks=[{"what": "fuga di gas in casa", "kind": "sicurezza"}])
    assert [t["id"] for t in fs.procedural_tools(gas, [])] == ["provvedimento_urgenza"]
    assert fs.procedural_tools(gas, ["penale"]) == []
    assert fs.procedural_tools(F(events=[ev("Tizio non paga il canone")]), ["contratti"]) == []
    assert fs.tool_refs({"tools": tools}) == [("696", "c.p.c."), ("696-bis", "c.p.c.")]


def test_extraction_asks_for_the_regime_fields_only_with_the_flag(monkeypatch):
    seen = []
    monkeypatch.setattr(fs, "_call_chat", lambda messages, **k: seen.append(messages[0].content) or json.dumps(
        {"parties": [{"name": "Anna", "role": "cliente", "kind": "persona fisica", "purpose": "personale"}]}))
    facts, _ = fs.extract_facts("Anna compra un divano.")
    assert "purpose" not in seen[-1] and set(facts["parties"][0]) == {"name", "role", "age"}
    monkeypatch.setenv(fs.REGIMES_FLAG, "true")
    facts, _ = fs.extract_facts("Anna compra un divano.")
    assert "purpose (per che cosa ha contratto" in seen[-1] and facts["parties"][0]["purpose"] == "personale"
    assert facts["risks"] == []


def test_regimes_alone_qualify_without_firing_the_rules(monkeypatch):
    monkeypatch.setenv(fs.REGIMES_FLAG, "true")
    monkeypatch.setenv("CASE_FRAME_ENABLED", "true")
    monkeypatch.setattr(dg, "_case_frame", lambda facts, request: {
        "tasks": ["Che tutele ha Anna?"], "areas": ["contratti"], "central_issue_candidates": []})
    monkeypatch.setattr(fs, "extract_facts", lambda text: (SOFA, ""))
    frame, signals = dg._frame_and_fact_signals("Anna compra un divano difettoso da Arredi Beta s.r.l.", "")
    assert signals["fired"] == [] and [q["quality"] for q in signals["qualifications"]] == ["consumatore", "professionista"]
    assert [c["nome"] for c in frame["central_issue_candidates"]] == ["garanzia legale di conformità", "foro del consumatore"]
    assert frame["central_issue_candidates"][0]["fonte"] == "qualifica consumatore_garanzia"
    assert "Anna: consumatore" in frame["qualifications"] and "Anna: consumatore" in dg._frame_block(frame)
    analysis = dg._with_central_questions({"questioni": []}, frame)
    assert analysis["questioni"][0]["istituti"][0]["ranges"] == [[128, 135]]


class _Session:
    def __init__(self, sections):
        self.sections, self.queries = sections, []

    def run(self, query, **params):
        self.queries.append(params)
        rows = [{"d": {"id": params["doc_id"]}, "s": s} for s in self.sections
                if params["phrase"] in s["plain_text"].lower()]
        return type("R", (), {"data": lambda self_: rows})()


def test_term_articles_found_by_their_words(monkeypatch):
    documents = [{"id": "CONS", "name": "Codice del Consumo"}]
    session = _Session([{"name": "133.0.0", "plain_text": "L'azione ... si prescrive nel termine di ventisei mesi "
                                                         "dalla consegna del bene."}])
    monkeypatch.setattr(dg, "_fetch_any", lambda s, d, refs: [
        {"_article": refs[0], "s": {"plain_text": "Il compratore decade dal diritto alla garanzia..."}}])
    facts, _ = fs.qualify_parties(SOFA)
    signals = {"terms": fs.compute_terms(facts, ["contratti"])}
    fs.attach_term_sources(session, documents, signals)
    by_id = {t["id"]: t for t in signals["terms"]}
    assert by_id["azione_conformita_consumo"]["ref"] == "art. 133 Codice del Consumo"
    assert by_id["azione_conformita_consumo"]["sources"][0]["text"].startswith("L'azione")
    assert by_id["denuncia_conformita_consumo"]["ref"] == "Codice del Consumo"        # not found: the code alone
    assert by_id["denuncia_vizi_vendita"]["sources"] == [{"ref": "art. 1495 c.c.",
                                                          "text": "Il compratore decade dal diritto alla garanzia..."}]


def test_procedural_tool_rows_join_after_the_check_and_are_protected(monkeypatch):
    existing = {"_article": ("700", "c.p.c."), "s": {"plain_text": "x"}}
    other = {"_article": ("1490", "c.c."), "s": {"plain_text": "y"}}
    added = {"_article": ("696", "c.p.c."), "_procedural_tool": "accertamento_preventivo", "s": {"plain_text": "z"}}
    tagged = {"_article": ("700", "c.p.c."), "_procedural_tool": "provvedimento_urgenza", "s": {"plain_text": "x"}}
    rows = dg._with_regime_rows([existing, other], [added, tagged])
    assert [r["_article"] for r in rows] == [("700", "c.p.c."), ("1490", "c.c."), ("696", "c.p.c.")]
    assert rows[0]["_procedural_tool"] == "provvedimento_urgenza"
    assert dg._protected_refs(rows, [])[("696", "c.p.c.")] == "procedural tool for the facts"


# --- the lawyer's review of the second case's answer: posture, criminal route, damages, roles --------------

def test_claimants_and_claims_received():
    claim = fs.claim_posture("Il figlio è rimasto ferito. I genitori vogliono ottenere il risarcimento dei danni.")
    assert claim["claimant"] and not claim["received"]
    received = fs.claim_posture("Il negozio ha ricevuto una lettera di diffida dal cliente. Come rispondere?")
    assert received["received"] and not received["claimant"]
    assert fs.claim_posture("Che cosa prevede la legge?") == {"claimant": False, "received": False, "why": ""}
    assert fs.claim_posture("Cosa possiamo fare?", F(requests=["possono chiedere il risarcimento del danno?"]))["claimant"]


def test_a_claimant_frame_drops_the_answer_a_claim_posture(monkeypatch):
    logged = []
    monkeypatch.setattr(dg, "vlog", lambda e, d=None, *a, **k: logged.append((e, d)))
    frame = {"procedural_posture": "Risposta alla richiesta di risarcimento ricevuta", "posture_type": "altro",
             "areas": ["responsabilita_civile"]}
    out = dg._claimant_frame(frame, {"claimant": True, "received": False, "why": "x"}, F())
    assert out["procedural_posture"].startswith("Nessun giudizio avviato e nessuna pretesa ricevuta")
    assert out["posture_type"] == "precontenzioso" and out["client_role"].startswith("attore")
    assert "Ruolo della parte assistita: attore" in dg._frame_block(out)
    assert dg._claimant_frame(frame, {"claimant": True, "received": True}, F()) == frame
    assert dg._claimant_frame(frame, {"claimant": False, "received": False}, F()) == frame


def test_the_checklist_follows_the_claim_posture(monkeypatch):
    monkeypatch.setenv(fs.REGIMES_FLAG, "true")
    from src.rag import practical_checklist
    checks = practical_checklist.POSTURE_CHECKS["precontenzioso"]
    analysis = {"fact_signals": {"posture": {"claimant": True, "received": False}}}
    out = ptw._by_claim_posture(checks, analysis)
    assert not any(c.startswith("Se la parte assistita riceve") for c in out)
    assert any(c.startswith("Messa in mora (art. 1219 c.c.)") for c in out)
    analysis = {"fact_signals": {"posture": {"claimant": False, "received": True}}}
    assert ptw._by_claim_posture(checks, analysis) == checks
    monkeypatch.delenv(fs.REGIMES_FLAG)
    assert ptw._by_claim_posture(checks, {"fact_signals": {"posture": {"claimant": True}}}) == checks


def test_a_personal_harm_by_negligence_brings_the_criminal_route():
    death = F(events=[ev("durante la gita il ragazzo cade dal pontile ed è morto")],
              injuries=[{"who": "il ragazzo", "days_to_heal": "", "permanent": "morte"}])
    tool = {t["id"]: t for t in fs.procedural_tools(death, ["responsabilita_civile"])}["profilo_penale"]
    assert tool["refs"] == [["589", "c.p."], ["76", "c.p.p."], ["2947", "c.c."]]
    assert "art. 2947, terzo comma, c.c." in tool["text"]
    injury = F(events=[ev("Tizio inciampa nella buca e si frattura il polso")],
               injuries=[{"who": "Tizio", "days_to_heal": "40", "permanent": "no"}])
    tool = {t["id"]: t for t in fs.procedural_tools(injury, ["responsabilita_civile"])}["profilo_penale"]
    assert ["590", "c.p."] in tool["refs"] and ["589", "c.p."] not in tool["refs"]
    beaten = F(events=[ev("Caio aggredisce Tizio e lo colpisce")], injuries=[{"who": "Tizio", "days_to_heal": "10"}])
    assert "profilo_penale" not in [t["id"] for t in fs.procedural_tools(beaten, [])]


def test_heads_of_damage_for_a_death_only():
    death = F(events=[ev("il paziente è deceduto il giorno dopo")])
    heads = fs.damage_heads(death)
    assert [h["id"] for h in heads] == ["danni_morte"] and ["2059", "c.c."] in heads[0]["refs"]
    block = fs.damages_block({"damages": heads})
    assert "Iure proprio" in block and "Iure hereditatis" in block and "art. 1917 c.c." in block
    assert fs.damage_heads(F(events=[ev("Tizio si frattura il polso")])) == []
    assert fs.regime_refs({"damages": heads, "tools": []}) == {("2059", "c.c."): "danni_morte",
                                                                ("1917", "c.c."): "danni_morte"}


def test_a_party_with_two_roles_is_examined_for_each():
    facts = F(parties=[person("Edil Gamma s.r.l.", kind="società", activities="ha costruito il capannone, ne cura la "
                              "manutenzione"), person("Anna", activities="ha acquistato")])
    roles = fs.multiple_roles(facts)
    assert roles == [{"party": "Edil Gamma s.r.l.", "activities": ["ha costruito il capannone", "ne cura la manutenzione"]}]
    assert "la responsabilità va valutata per ciascun ruolo" in fs.roles_block({"multi_roles": roles})


def test_blind_answers_lose_the_deep_dive_suggestion(tmp_path):
    from docx import Document
    docx_script = _script("answer_docx")
    answer = ("**6. Conclusioni**\n\nIl rischio è contenuto, vedi [custodia](deepdive:custodia).\n\n"
              "Vuoi approfondire in particolare [custodia](deepdive:custodia) o [caso fortuito](deepdive:caso-fortuito)?"
              "\n\n⚠️ Verificare tutto prima dell'uso.")
    assert docx_script.without_deep_dive(answer) == ("**6. Conclusioni**\n\nIl rischio è contenuto, vedi custodia."
                                                     "\n\n⚠️ Verificare tutto prima dell'uso.")
    out = str(tmp_path / "r.docx")
    docx_script.build({"status": "ok", "answer": answer}, out, "Risposta 1", "", None, blind=True)
    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "approfondire" not in text and "deepdive" not in text
    docx_script.build({"status": "ok", "answer": answer}, out, "Risposta", "", None)
    assert "Vuoi approfondire" in "\n".join(p.text for p in Document(out).paragraphs)
